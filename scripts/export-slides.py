"""Render a Quarto Reveal deck to PNG slides and an image-based PowerPoint."""

from __future__ import annotations

import argparse
from contextlib import contextmanager
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import threading
from urllib.parse import quote

ROOT = Path(__file__).resolve().parents[1]
# Keep the browser installation local; callers may override this variable.
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", str(ROOT / "tools/playwright"))


class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, format, *args):
        pass


@contextmanager
def serve(directory):
    server = ThreadingHTTPServer(
        ("127.0.0.1", 0), partial(QuietHandler, directory=str(directory))
    )
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}"
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


READY_IMAGES = """async () => {
    await document.fonts.ready;
    const slide = Reveal.getCurrentSlide();
    const background = Reveal.getSlideBackground();
    const roots = [slide, background].filter(Boolean);
    const images = roots.flatMap(root => Array.from(root.querySelectorAll('img')));
    await Promise.all(images.map(async img => {
        if (img.dataset.src && !img.getAttribute('src')) img.src = img.dataset.src;
        try { await img.decode(); }
        catch { throw new Error('Image failed to load: ' + (img.alt || img.src.slice(0, 160))); }
        if (!img.naturalWidth) throw new Error('Empty image: ' + img.alt);
    }));
    // Background images are CSS resources, not <img> elements.
    const urls = new Set();
    for (const root of roots) {
        for (const el of [root, ...root.querySelectorAll('*')]) {
            const css = getComputedStyle(el).backgroundImage;
            for (const match of css.matchAll(/url\\(["']?(.*?)["']?\\)/g)) urls.add(match[1]);
        }
    }
    await Promise.all(Array.from(urls, async url => {
        const image = new Image(); image.src = url;
        try { await image.decode(); }
        catch { throw new Error('Background failed to load: ' + url.slice(0, 160)); }
    }));
    if ([...document.fonts].some(font => font.status === 'error')) {
        throw new Error('A presentation font failed to load');
    }
    await new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve)));
}"""


def export(args):
    try:
        from playwright.sync_api import sync_playwright
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise RuntimeError(
            "Missing export dependencies. Run: python -m pip install -r "
            "scripts/requirements-export.txt"
        ) from exc

    source = Path(args.input).resolve()
    if not source.is_file() or source.suffix.lower() != ".qmd":
        raise ValueError(f"Expected an existing .qmd deck: {source}")
    if not source.is_relative_to(ROOT):
        raise ValueError("The deck must be inside this Quarto repository.")
    output = (ROOT / "export" / source.stem / "image-pptx").resolve()
    html_dir = output / "html"
    images_dir = output / "images"
    images_dir.mkdir(parents=True, exist_ok=True)
    quarto = shutil.which("quarto")
    if not quarto:
        raise RuntimeError("Quarto is not on PATH.")

    with sync_playwright() as playwright:
        if args.install_browser:
            subprocess.run(
                [sys.executable, "-m", "playwright", "install", "chromium"], check=True
            )
        if not Path(playwright.chromium.executable_path).is_file():
            raise RuntimeError("Chromium is missing. Run this script with --install-browser once.")
        print(f"Rendering {source.name}...", flush=True)
        subprocess.run(
            [quarto, "render", str(source.relative_to(ROOT)), "--profile", "export",
             "--to", "revealjs", "--output-dir", str(html_dir.relative_to(ROOT))],
            cwd=ROOT, check=True,
        )
        html = html_dir / source.relative_to(ROOT).with_suffix(".html")
        if not html.is_file():
            raise RuntimeError(f"Quarto did not create the expected HTML: {html}")

        with serve(html_dir) as url:
            browser = playwright.chromium.launch(headless=True)
            try:
                page = browser.new_page(
                    viewport={"width": 1920, "height": 1080},
                    device_scale_factor=args.scale,
                )
                page.set_default_timeout(60000)
                page.goto(url + "/" + quote(html.relative_to(html_dir).as_posix()))
                page.wait_for_function("window.Reveal && Reveal.isReady()")
                config = page.evaluate("({width: Reveal.getConfig().width, height: Reveal.getConfig().height})")
                width, height = config["width"], config["height"]
                if not all(isinstance(v, (int, float)) and v > 0 for v in (width, height)):
                    raise ValueError("Reveal width and height must be numeric pixel dimensions.")
                width, height = int(width), int(height)
                page.set_viewport_size({"width": width, "height": height})
                page.evaluate("""() => {
                    Reveal.configure({transition: 'none', backgroundTransition: 'none',
                        autoSlide: 0, autoAnimate: false, controls: false, progress: false,
                        fragments: false, view: 'slide', scrollActivationWidth: null});
                    Reveal.layout();
                }""")
                page.add_style_tag(content="""
                    .slide-menu-button, .slide-chalkboard-buttons, .reveal .playback,
                    .reveal .slide-menu, .reveal .slide-menu-overlay { display: none !important; }
                    *, *::before, *::after { transition: none !important; }
                """)
                slides = page.evaluate("""() => Reveal.getSlides().map(slide => ({
                    ...Reveal.getIndices(slide),
                    title: slide.querySelector('h1,h2,h3')?.textContent || '',
                    notes: Reveal.getSlideNotes(slide) || ''
                }))""")
                if not slides:
                    raise RuntimeError("No Reveal slides found.")
                presentation = Presentation()
                presentation.slide_width = Inches(13.333333)
                presentation.slide_height = round(presentation.slide_width * height / width)
                presentation.core_properties.title = page.title()
                records = []
                for number, info in enumerate(slides, 1):
                    page.evaluate("s => Reveal.slide(s.h, s.v || 0)", info)
                    page.evaluate("""async ready => {
                        let timer;
                        try {
                            await Promise.race([
                                (new Function('return (' + ready + ')()'))(),
                                new Promise((_, reject) => {
                                    timer = setTimeout(() => reject(new Error(
                                        'Timed out waiting for slide images/fonts')), 45000);
                                })
                            ]);
                        } finally { clearTimeout(timer); }
                    }""", READY_IMAGES)
                    # Give slide-change handlers a chance to settle after image decoding.
                    page.wait_for_timeout(150)
                    image = images_dir / f"slide-{number:03d}.png"
                    page.screenshot(path=str(image), animations="disabled", scale="device")
                    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                    slide.shapes.add_picture(
                        str(image), 0, 0,
                        width=presentation.slide_width, height=presentation.slide_height,
                    )
                    if info["notes"]:
                        notes = page.evaluate("""html => {
                            const div = document.createElement('div'); div.innerHTML = html;
                            return div.textContent;
                        }""", info["notes"])
                        slide.notes_slide.notes_text_frame.text = notes
                    records.append({"number": number, "title": info["title"], "image": image.name})
                    print(f"  {number}/{len(slides)} {info['title']}", flush=True)
                pptx = output / f"{source.stem}-images.pptx"
                temporary = output / f"{source.stem}-images.tmp.pptx"
                presentation.save(temporary)
                temporary.replace(pptx)
                (output / "manifest.json").write_text(json.dumps({
                    "source": source.relative_to(ROOT).as_posix(),
                    "image_width": round(width * args.scale),
                    "image_height": round(height * args.scale), "slides": records,
                }, indent=2), encoding="utf-8")
                print(f"\nPowerPoint: {pptx}\nSlide images: {images_dir}", flush=True)
            finally:
                browser.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", help="Path to a .qmd deck in this repository")
    parser.add_argument("--scale", type=int, choices=range(1, 5), default=2,
                        help="Pixels per CSS pixel (default: 2; 3840x2160 for this theme)")
    parser.add_argument("--install-browser", action="store_true",
                        help="Download the Playwright Chromium browser before exporting")
    args = parser.parse_args()
    try:
        export(args)
    except Exception as exc:
        print(f"Export failed: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
